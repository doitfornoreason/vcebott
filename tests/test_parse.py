from pathlib import Path

from docx import Document

from vcebott.parse import _match_rows, parse_report_docx, tidy_table


def _match(marks_tokens, pct_tokens, has_average=True):
    return _match_rows(marks_tokens, pct_tokens, has_average)


def test_match_rows_basic():
    parsed = _match(["0", "1", "Average"], ["42", "58", "0.6"])
    assert parsed is not None
    marks_available, dist, average = parsed
    assert marks_available == 1
    assert dist == {0: 42.0, 1: 58.0}
    assert average == pytest_approx(0.6)


def pytest_approx(value):
    import pytest

    return pytest.approx(value)


def test_match_rows_dash_zero():
    parsed = _match(["0", "1", "2", "Average"], ["10", "--", "90", "1.4"])
    assert parsed is not None
    _, dist, average = parsed
    assert dist == {0: 10.0, 1: 0.0, 2: 90.0}
    assert average == pytest_approx(1.4)


def test_match_rows_no_average_column():
    parsed = _match(["0", "1"], ["42", "58"], has_average=False)
    assert parsed is not None
    _, dist, average = parsed
    assert dist == {0: 42.0, 1: 58.0}
    assert average is None


def test_match_rows_mismatch_rejected():
    assert _match(["0", "1", "2"], ["42", "58"], has_average=False) is None


def _add_marks_table(doc, label, marks, pcts, average=None):
    doc.add_paragraph(f"Question {label}")
    table = doc.add_table(rows=2, cols=len(marks) + (2 if average is not None else 1))
    header = ["Marks", *map(str, marks)]
    values = ["%", *map(str, pcts)]
    if average is not None:
        header.append("Average")
        values.append(str(average))
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = header[i]
    for i, cell in enumerate(table.rows[1].cells):
        cell.text = values[i]


def test_parse_report_docx(tmp_path: Path):
    doc = Document()
    _add_marks_table(doc, "1a.", [0, 1], [42, 58], average=0.6)
    doc.add_paragraph("Some commentary text.")
    _add_marks_table(doc, "1b.", [0, 1, 2], [10, 39, 51], average=1.4)
    _add_marks_table(doc, "2", [0, 1], ["5", "95"], average="0.9")
    path = tmp_path / "report.docx"
    doc.save(str(path))

    stats = parse_report_docx(path)
    assert [s.question for s in stats] == ["1a", "1b", "2"]
    assert stats[0].marks_available == 1 and stats[0].average == pytest_approx(0.6)
    assert stats[1].distribution == {0: 10.0, 1: 39.0, 2: 51.0}
    assert stats[2].distribution == {0: 5.0, 1: 95.0}

    frame = tidy_table("methods", 2023, 1, False, stats)
    assert list(frame.columns) == [
        "subject",
        "year",
        "exam",
        "nht",
        "question",
        "marks_available",
        "mark",
        "pct",
        "average",
    ]
    assert len(frame) == 2 + 3 + 2
    assert frame.iloc[0]["question"] == "1a"
    assert frame.iloc[0]["mark"] == 0 and frame.iloc[0]["pct"] == 42.0
